"""
generate_full_report.py
-----------------------
Generates a comprehensive Word document covering:
  - The full YOLO-Seg pipeline
  - Training results
  - External test set evaluation (534 images)
  - Comparison with all previous models
  - Dish-level breakdown
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "yolo_seg_pipeline_report.docx")
CURVES_IMG  = os.path.abspath("results_mobilenet_v3_large_yolo_seg/training_curves.png")
CM_IMG      = os.path.abspath("results/test_external_mobilenet_v3_large_yolo_seg/confusion_matrix.png")

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.upper())
    tcPr.append(shd)

def add_borders(table):
    for tr in table._tbl.findall(qn("w:tr")):
        for tc in tr.findall(qn("w:tc")):
            tcPr = tc.get_or_add_tcPr()
            tcB  = OxmlElement("w:tcBorders")
            for side in ("top","left","bottom","right","insideH","insideV"):
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"),   "single")
                b.set(qn("w:sz"),    "4")
                b.set(qn("w:space"), "0")
                b.set(qn("w:color"), "CCCCCC")
                tcB.append(b)
            tcPr.append(tcB)

def heading(doc, text, level=1, color="1F3864"):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def para(doc, text="", bold=False, italic=False, size=11, color=None):
    p   = doc.add_paragraph()
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def bullet(doc, text, prefix=None, prefix_color="1F3864"):
    p = doc.add_paragraph(style="List Bullet")
    if prefix:
        r = p.add_run(prefix)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(prefix_color)
    p.add_run(text)
    return p

def make_table(doc, headers, rows, hdr_bg="1F3864", alt_bg="EBF3FB", col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bg(cell, hdr_bg)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # rows
    for ri, row_data in enumerate(rows):
        bg  = alt_bg if ri % 2 else "FFFFFF"
        row = table.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = str(val)
            set_cell_bg(cell, bg)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)

    add_borders(table)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table

def insert_image(doc, img_path, width_cm=14, caption=None):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Cm(width_cm))
        if caption:
            cp = doc.add_paragraph(caption)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cp.runs:
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    else:
        doc.add_paragraph(f"[Image not found: {img_path}]").italic = True


# ── Build Document ─────────────────────────────────────────────────────────────

doc = Document()

# margins
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin  = sec.bottom_margin = Cm(2.0)

# ── Cover ─────────────────────────────────────────────────────────────────────
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("🍽️  Food Classifier — YOLO-Seg Pipeline Report")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor.from_string("1F3864")

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y  %H:%M')}").italic = True

doc.add_paragraph()

# ── 1. Executive Summary ──────────────────────────────────────────────────────
heading(doc, "1. Executive Summary", 1)

p = doc.add_paragraph()
p.add_run(
    "This report presents the results of the upgraded food classification pipeline. "
    "The key improvement over previous runs is the use of "
)
p.add_run("YOLOv8 Segmentation").bold = True
p.add_run(
    " (mask-based tight cropping) for pre-processing, combined with "
    "MobileNet V3 Large transfer learning. "
    "The model was trained on one dataset and evaluated on a fully "
)
p.add_run("separate, held-out external test set").bold = True
p.add_run(
    " of 534 images — giving the most honest accuracy measurement yet."
)

doc.add_paragraph()

# KPI boxes via a 3-col table
make_table(doc,
    headers=["🎯 External Test Accuracy", "⚡ Val Accuracy", "📦 Test Images"],
    rows=[["92.70%", "94.58%", "534 (26 dishes)"]],
    hdr_bg="2E75B6", alt_bg="DEEAF1",
)
doc.add_paragraph()

# ── 2. Pipeline Overview ──────────────────────────────────────────────────────
heading(doc, "2. Pipeline Overview", 1)

make_table(doc,
    headers=["Step", "Action", "Output"],
    rows=[
        ["1", "YOLO-Seg crop — training data",  "results/yolo_seg_dataset  (1,203 imgs)"],
        ["2", "YOLO-Seg crop — test data",       "results/yolo_seg_test     (534 imgs)"],
        ["3", "Train MobileNet V3 Large",        "checkpoints/mobilenet_v3_large_yolo_seg_best.pth"],
        ["4", "Evaluate on external test set",   "results/test_external_mobilenet_v3_large_yolo_seg/"],
    ],
    col_widths=[1.2, 6.0, 7.5],
)
doc.add_paragraph()

p = doc.add_paragraph()
p.add_run("Why YOLO Segmentation?  ").bold = True
p.add_run(
    "Unlike detection (bounding box), segmentation extracts a polygon mask around "
    "the food item, yielding a tighter crop with less background noise. "
    "In this run, "
)
p.add_run("93.9% of training images (1,130/1,203) received a proper seg-mask crop").bold = True
p.add_run(", vs only a center-crop fallback for the remaining 6.1%.")

doc.add_paragraph()

# ── 3. YOLO-Seg Cropping Stats ────────────────────────────────────────────────
heading(doc, "3. YOLO-Seg Cropping Statistics", 1)

make_table(doc,
    headers=["Dataset", "Total", "Seg-Mask ✅", "Center-Crop ℹ️", "Skipped ⚠️"],
    rows=[
        ["Training (categorized_dataset)", "1,204", "1,130  (93.9%)", "73  (6.1%)", "1"],
        ["Test (test_dataset_dishes)",     "534",   "481  (90.1%)",   "53  (9.9%)", "0"],
    ],
)
doc.add_paragraph()

# ── 4. Training Configuration ─────────────────────────────────────────────────
heading(doc, "4. Training Configuration", 1)

make_table(doc,
    headers=["Parameter", "Value"],
    rows=[
        ["Architecture",       "MobileNet V3 Large (ImageNet pretrained)"],
        ["Dataset",            "results/yolo_seg_dataset  (1,203 images)"],
        ["Train / Val split",  "80% train  /  20% val  (no internal test split)"],
        ["Image size",         "288 × 288 px"],
        ["Batch size",         "32"],
        ["Stage 1 (frozen)",   "15 epochs — classifier head only, LR=1e-3, ReduceLROnPlateau"],
        ["Stage 2 (unfrozen)", "25 epochs — full fine-tune, LR=1e-5, CosineAnnealingLR"],
        ["Loss function",      "CrossEntropyLoss  (label smoothing=0.1 + class weights)"],
        ["Early stopping",     "Patience = 7 epochs (on val loss)"],
        ["Total epochs run",   "40  (15 + 25, no early stop triggered)"],
        ["Device",             "Apple MPS (M-series GPU)"],
    ],
    col_widths=[5.5, 10.0],
)
doc.add_paragraph()

# ── 5. Training Results ───────────────────────────────────────────────────────
heading(doc, "5. Training Results", 1)

make_table(doc,
    headers=["Metric", "Stage 1 End", "Stage 2 End (Best)"],
    rows=[
        ["Train Accuracy", "~85%",   "99.07%"],
        ["Val Accuracy",   "~88%",   "94.58%  ★"],
        ["Train Loss",     "~0.45",  "0.2426"],
        ["Val Loss",       "~0.38",  "0.3357"],
        ["Train/Val Gap",  "—",      "~4.5%  (low overfitting ✅)"],
    ],
)
doc.add_paragraph()

insert_image(doc, CURVES_IMG, width_cm=14,
             caption="Figure 1 — Training & Validation curves across 40 epochs (Stage 1 + Stage 2)")
doc.add_paragraph()

# ── 6. External Test Set Results ─────────────────────────────────────────────
heading(doc, "6. External Test Set Results", 1)

p = doc.add_paragraph()
p.add_run("Test set: ").bold = True
p.add_run(
    "534 images across 26 dishes (11 non_veg + 15 veg), "
    "completely separate from the training dataset — "
    "added manually by the user and never seen during training or validation."
)
doc.add_paragraph()

heading(doc, "6.1  Overall Metrics", 2, color="2E75B6")
make_table(doc,
    headers=["Metric", "Score"],
    rows=[
        ["Accuracy",              "92.70%  (495 / 534 correct)"],
        ["Precision (weighted)",  "92.70%"],
        ["Recall (weighted)",     "92.70%"],
        ["F1-Score (weighted)",   "92.70%"],
        ["Incorrect predictions", "39 / 534"],
    ],
    col_widths=[6.0, 9.5],
)
doc.add_paragraph()

heading(doc, "6.2  Per-Class Results", 2, color="2E75B6")
make_table(doc,
    headers=["Class", "Precision", "Recall", "F1-Score", "Support"],
    rows=[
        ["non_veg", "92.96%", "92.62%", "92.79%", "271"],
        ["veg",     "92.42%", "92.78%", "92.60%", "263"],
        ["Overall (weighted)", "92.70%", "92.70%", "92.70%", "534"],
    ],
)
doc.add_paragraph()

heading(doc, "6.3  Confusion Matrix", 2, color="2E75B6")
make_table(doc,
    headers=["", "Predicted: non_veg", "Predicted: veg"],
    rows=[
        ["Actual: non_veg", "TN = 251  ✅", "FP = 20  ❌"],
        ["Actual: veg",     "FN = 19  ❌",  "TP = 244  ✅"],
    ],
    col_widths=[4.5, 5.5, 5.5],
)
doc.add_paragraph()

insert_image(doc, CM_IMG, width_cm=10,
             caption="Figure 2 — Confusion matrix on external test set (534 images)")
doc.add_paragraph()

# ── 7. Dish-Level Breakdown ───────────────────────────────────────────────────
heading(doc, "7. Test Set — Dish-Level Breakdown", 1)

make_table(doc,
    headers=["Category", "Dish", "Images in Test Set"],
    rows=[
        # non_veg
        ["non_veg", "grilled_kofta",    "33"],
        ["non_veg", "kofta_stew",       "23"],
        ["non_veg", "mandi",            "26"],
        ["non_veg", "meat-madfoon",     "24"],
        ["non_veg", "meat_maqluba",     "31"],
        ["non_veg", "meat_pie",         "29"],
        ["non_veg", "mixed_grill",      "19"],
        ["non_veg", "shawarma",         "22"],
        ["non_veg", "shish tawook",     "16"],
        ["non_veg", "shrimp_sayadiyah", "23"],
        ["non_veg", "stuffed_pigeon",   "25"],
        # veg
        ["veg", "cheese_pie",          "23"],
        ["veg", "hummus",              "21"],
        ["veg", "okra_tagine",         "15"],
        ["veg", "onion_soup",          "19"],
        ["veg", "pistachio_sweet",     "17"],
        ["veg", "pistachio_tamriya",   "12"],
        ["veg", "pomegranate_fattoush","16"],
        ["veg", "rice_pudding",        "19"],
        ["veg", "saffron_cake",        "22"],
        ["veg", "sesame_brittle",      "19"],
        ["veg", "spinach_fatayer",     "14"],
        ["veg", "spinach_stew",        "16"],
        ["veg", "tamees_bread",        "18"],
        ["veg", "tamriya",             "12"],
        ["veg", "za'atar_manakish",    "20"],
    ],
    col_widths=[3.0, 6.5, 4.5],
)
doc.add_paragraph()

# ── 8. Model Comparison ───────────────────────────────────────────────────────
heading(doc, "8. Comparison With Previous Models", 1)

make_table(doc,
    headers=["Model", "Test Accuracy", "Test Set", "Overfitting?"],
    rows=[
        ["MobileNet V3 + YOLO detect (old)",  "81.7%  (val)",        "Same dataset — 15% split",    "❌ High (~18% gap)"],
        ["EfficientNet B2 (no YOLO)",          "89.3%",               "Same dataset — 15% split",    "✅ Low"],
        ["MobileNet V3 Large (no YOLO)",       "92.2%",               "Same dataset — 15% split",    "⚠️ Moderate"],
        ["MobileNet V3 + YOLO-Seg  ★ NEW",    "92.70%",              "External — 534 fresh images", "✅ Low (4.5% gap)"],
    ],
    hdr_bg="1F3864",
)
doc.add_paragraph()

p = doc.add_paragraph()
p.add_run("Key insight: ").bold = True
p.add_run(
    "The new YOLO-Seg model achieves 92.70% on a completely external test set "
    "of 534 fresh images — the most reliable accuracy number yet. "
    "The train/val gap of only ~4.5% confirms that overfitting has been significantly reduced "
    "compared to the old YOLO-detect pipeline (~18% gap)."
)
doc.add_paragraph()

# ── 9. Next Steps ─────────────────────────────────────────────────────────────
heading(doc, "9. Recommended Next Steps", 1)

make_table(doc,
    headers=["Priority", "Action", "Expected Gain"],
    rows=[
        ["🔴 High",   "Add 10+ high-quality images to weak dishes (tamriya=12, pistachio_tamriya=12)", "+1–3%"],
        ["🔴 High",   "MixUp / CutMix augmentation in training pipeline",                               "+2–4%"],
        ["🟠 Medium", "Test-Time Augmentation (TTA) at inference",                                      "+1–2%"],
        ["🟠 Medium", "Ensemble MobileNet V3 + EfficientNet B2",                                        "+2–3%"],
        ["🟡 Low",    "Try EfficientNet B4 backbone (larger model)",                                    "+2–5%"],
    ],
    col_widths=[2.5, 9.0, 3.5],
)
doc.add_paragraph()

# ── Footer ────────────────────────────────────────────────────────────────────
fp = doc.add_paragraph()
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(f"Food Classifier Project  •  YOLO-Seg Pipeline Report  •  {datetime.now().strftime('%Y-%m-%d')}")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(OUTPUT_PATH)
print(f"✅  Report saved to: {OUTPUT_PATH}")
