import sys
import os
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import os
import json
import torch
from torchvision import transforms
from torch.utils.data import DataLoader, Dataset
from PIL import Image
from tqdm import tqdm
from core.models import FoodClassifier

class InferenceDataset(Dataset):
    def __init__(self, file_paths, transform):
        self.file_paths = file_paths
        self.transform = transform
        
    def __len__(self):
        return len(self.file_paths)
        
    def __getitem__(self, idx):
        path = self.file_paths[idx]
        image = Image.open(path).convert('RGB')
        return self.transform(image), path

def get_transforms(img_size):
    return transforms.Compose([
        transforms.Resize((img_size, img_size)),
        transforms.ToTensor(),
        transforms.Normalize(mean=[0.485, 0.456, 0.406], std=[0.229, 0.224, 0.225])
    ])

def main():
    device = torch.device('mps' if torch.backends.mps.is_available() else 'cpu')
    
    # Gather all valid images from split_info
    with open('results_mobilenet_v3_large/split_info.json', 'r') as f:
        split_info = json.load(f)
        all_files = split_info.get('train_files', []) + split_info.get('val_files', [])
        
    print(f"Found {len(all_files)} total valid images in the dataset.")
    
    models = [
        ('mobilenet_v3_large', 224),
        ('efficientnet_b2', 288)
    ]
    
    results = {}
    
    for model_name, img_size in models:
        print(f"\nEvaluating {model_name} on entire dataset...")
        transform = get_transforms(img_size)
        dataset = InferenceDataset(all_files, transform)
        loader = DataLoader(dataset, batch_size=32, shuffle=False, num_workers=4)
        
        # Load mapping
        with open(f'results_{model_name}/class_mapping.json', 'r') as f:
            mapping = json.load(f)
            
        model = FoodClassifier(num_classes=2, freeze_backbone=False, architecture=model_name)
        model.load_state_dict(torch.load(f'checkpoints/{model_name}_best.pth', map_location=device))
        model.to(device)
        model.eval()
        
        veg_pred = 0
        non_veg_pred = 0
        
        import shutil
        out_dir = f'results_{model_name}/total_predictions'
        os.makedirs(os.path.join(out_dir, 'veg'), exist_ok=True)
        os.makedirs(os.path.join(out_dir, 'non_veg'), exist_ok=True)
        
        with torch.no_grad():
            for inputs, paths in tqdm(loader, desc=f"Predicting with {model_name}"):
                inputs = inputs.to(device)
                outputs = model(inputs)
                _, preds = torch.max(outputs, 1)
                
                for p, fpath in zip(preds.cpu().numpy(), paths):
                    pred_class = mapping[str(p)]
                    if pred_class == 'veg':
                        veg_pred += 1
                    else:
                        non_veg_pred += 1
                        
                    basename = os.path.basename(fpath)
                    dish_name = os.path.basename(os.path.dirname(fpath))
                    dest_dir = os.path.join(out_dir, pred_class, dish_name)
                    os.makedirs(dest_dir, exist_ok=True)
                    shutil.copy(fpath, os.path.join(dest_dir, basename))
                        
        results[model_name] = {'veg': veg_pred, 'non_veg': non_veg_pred}
        print(f"{model_name} -> Predicted Veg: {veg_pred}, Predicted Non-Veg: {non_veg_pred}")
        
    # Update word document
    print("\nUpdating Word Document...")
    from docx import Document
    doc = Document('Food_Classifier_Project_Report.docx')
    
    # We will just append a new section at the end
    doc.add_heading('8. Total Dataset Predictions', 1)
    doc.add_paragraph(f"To see how each model interprets the entire dataset of {len(all_files)} images (including training data), we ran full inference on all images:")
    for model_name, res in results.items():
        doc.add_paragraph(f"{model_name}:", style='List Bullet')
        doc.add_paragraph(f"    - Total Predicted Vegetarian: {res['veg']}")
        doc.add_paragraph(f"    - Total Predicted Non-Vegetarian: {res['non_veg']}")
        
    doc.save('Food_Classifier_Project_Report.docx')
    print("Report updated!")

if __name__ == '__main__':
    main()
