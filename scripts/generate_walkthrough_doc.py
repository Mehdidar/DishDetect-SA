import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs", "Project_Walkthrough.docx")
ROOT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")

def heading(doc, text, level=1, color="1F3864"):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def main():
    doc = Document()
    
    # Adjust margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("Food Classifier Project\nDeveloper Walkthrough & Architecture")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("1F3864")

    doc.add_paragraph()
    
    # Introduction
    heading(doc, "1. Project Overview")
    doc.add_paragraph("This document serves as a developer walkthrough for the newly reorganized Food Classifier project. The monolithic structure has been modularized into distinct directories for models, scripts, tools, and UI applications.")

    # Project Structure
    heading(doc, "2. Codebase Structure", level=1)
    doc.add_paragraph("The root directory is clean and organized into logical components:")
    
    ul_c1 = doc.add_paragraph(style='List Bullet')
    ul_c1.add_run("core/: ").bold = True
    ul_c1.add_run("Contains the main application logic and model definitions. Inside core/models/, you will find dedicated model files like mobilenet.py and efficientnet.py. The __init__.py acts as a factory that loads the requested model.")
    
    ul_c2 = doc.add_paragraph(style='List Bullet')
    ul_c2.add_run("scripts/: ").bold = True
    ul_c2.add_run("Contains the active pipeline scripts used for model training and core evaluations (e.g., train.py, evaluate.py, yolo_seg_crop.py, predict.py).")
    
    ul_c3 = doc.add_paragraph(style='List Bullet')
    ul_c3.add_run("tools/: ").bold = True
    ul_c3.add_run("Contains all diagnostic and reporting scripts (e.g., dataset_check.py, full_dataset_diagnostic.py). Use these to check dataset integrity.")
    
    ul_c4 = doc.add_paragraph(style='List Bullet')
    ul_c4.add_run("app/: ").bold = True
    ul_c4.add_run("Contains your user interface scripts (ui.py, patch_ui.py) for running the interactive front-end.")
    
    ul_c5 = doc.add_paragraph(style='List Bullet')
    ul_c5.add_run("docs/: ").bold = True
    ul_c5.add_run("Contains generated project reports (e.g., Project_Evolution_History.docx, Architecture_Comparison_Report.docx).")

    # Common Workflows
    heading(doc, "3. Common Workflows & Commands", level=1)
    
    heading(doc, "Setup & Requirements", level=2)
    doc.add_paragraph("Install all project dependencies:")
    p = doc.add_paragraph()
    p.add_run("pip install -r requirements.txt").font.name = 'Courier New'

    heading(doc, "Validating the Dataset", level=2)
    doc.add_paragraph("Before training, always ensure the dataset is clean and valid. This generates a dataset_report.json:")
    p = doc.add_paragraph()
    p.add_run("python tools/dataset_check.py --dataset /path/to/dataset").font.name = 'Courier New'

    heading(doc, "Training the Model", level=2)
    doc.add_paragraph("Run the main training pipeline. By default, it will load the model from core.models based on your arguments. Training results are saved to results/ and checkpoints/:")
    p = doc.add_paragraph()
    p.add_run("python scripts/train.py --dataset /path/to/dataset").font.name = 'Courier New'

    heading(doc, "Evaluating the Model", level=2)
    doc.add_paragraph("After training, evaluate the model on the test split. This generates classification reports and a confusion matrix:")
    p = doc.add_paragraph()
    p.add_run("python scripts/evaluate.py --dataset /path/to/dataset").font.name = 'Courier New'

    heading(doc, "Running the UI Application", level=2)
    doc.add_paragraph("To start the graphical user interface for live predictions:")
    p = doc.add_paragraph()
    p.add_run("python app/ui.py").font.name = 'Courier New'

    # Performance Visualizations
    heading(doc, "4. Current Model Performance (MobileNetV3 + YOLO Seg)", level=1)
    doc.add_paragraph("The current primary architecture leverages a MobileNetV3-Large backbone with YOLO segmentation as a preprocessing step. Below are the key visualizations outlining the model's robustness.")
    
    heading(doc, "Training Convergence", level=2)
    doc.add_paragraph("The training and validation curves demonstrating convergence over the epochs:")
    curve_path = os.path.join(ROOT_DIR, "results_mobilenet_v3_large_yolo_seg", "training_curves.png")
    if os.path.exists(curve_path):
        doc.add_picture(curve_path, width=Inches(6.0))
    else:
        doc.add_paragraph("[Training Curves image not found]").italic = True

    heading(doc, "Standard Validation Results", level=2)
    doc.add_paragraph("Confusion matrix for the standard validation set, distinguishing between purely 'Veg' and 'Non-Veg' categories:")
    cm_path = os.path.join(ROOT_DIR, "results_mobilenet_v3_large_yolo_seg", "confusion_matrix.png")
    if os.path.exists(cm_path):
        doc.add_picture(cm_path, width=Inches(5.5))
    else:
        doc.add_paragraph("[Confusion Matrix image not found]").italic = True

    heading(doc, "Real-World Restaurant Testing", level=2)
    doc.add_paragraph("Testing against highly ambiguous out-of-distribution real-world restaurant images (e.g., hidden meat under rice in Mandi, mixed ingredient stews).")
    
    pie_path = os.path.join(ROOT_DIR, "results", "restaurant_analysis", "overall_accuracy_pie.png")
    if os.path.exists(pie_path):
        doc.add_picture(pie_path, width=Inches(5.0))
    else:
        doc.add_paragraph("[Overall Accuracy Pie image not found]").italic = True
        
    doc.add_paragraph("Per-Dish breakdown of the restaurant test:")
    dish_path = os.path.join(ROOT_DIR, "results", "restaurant_analysis", "per_dish_accuracy.png")
    if os.path.exists(dish_path):
        doc.add_picture(dish_path, width=Inches(6.0))
    else:
        doc.add_paragraph("[Per-Dish Accuracy image not found]").italic = True

    # Generating the Document
    doc.save(OUTPUT_PATH)
    print(f"Generated walkthrough document at: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
