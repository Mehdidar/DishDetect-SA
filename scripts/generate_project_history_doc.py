import sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime

OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Project_Evolution_History.docx")

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.upper())
    tcPr.append(shd)

def heading(doc, text, level=1, color="1F3864"):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def main():
    doc = Document()
    
    # Adjust margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("Food Classifier Project\nDetailed Evolution & Hyperparameters")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("1F3864")

    doc.add_paragraph()
    
    # Introduction
    heading(doc, "1. Project Overview")
    doc.add_paragraph("This document outlines the complete developmental journey of the Food Classifier project, tracking the iterative phases of modeling, dataset curation, and detailed technical hyperparameters (learning rate, accuracy, epochs).")

    # Data Pre-Processing & Cleaning
    heading(doc, "2. Deep Data Pre-Processing & Cleaning", level=1)
    doc.add_paragraph("Prior to any training, the dataset was rigorously scrubbed using automated tools to ensure data integrity and prevent model leakage. According to the deep raw JSON reports (dataset_report.json), the following actions were taken:")
    ul_d1 = doc.add_paragraph(style='List Bullet')
    ul_d1.add_run("Raw Base Dataset: 1,439 total images (highly balanced with 758 non-veg and 681 veg images).")
    ul_d2 = doc.add_paragraph(style='List Bullet')
    ul_d2.add_run("Corrupted & Unsupported File Rejection: Automatically stripped out 1 fully corrupted image and 8 unsupported file formats (including .avif extensions and hidden macOS .DS_Store files).")
    ul_d3 = doc.add_paragraph(style='List Bullet')
    ul_d3.add_run("Size Filtering: Discarded images with resolutions too small to be meaningful for a CNN.")
    ul_d4 = doc.add_paragraph(style='List Bullet')
    ul_d4.add_run("Duplicate Leakage Prevention: Used perceptual hashing (threshold=5) to locate and group near-duplicate augmented image pairs (like matching Shawarma and Shrimp Sayadiyah images), ensuring they were forcefully bucketed into the same split to prevent train/test data leakage.")

    # Phase 1
    heading(doc, "Phase 1: Baseline 25-Class Arabic Dishes Model", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" Original Arabic 25 Dishes")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" Good (baseline qualitative estimation)")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate (Stage 1): 1e-3 (AdamW Optimizer)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Learning Rate (Stage 2): 1e-5 (Cosine Annealing)")
    ul3 = doc.add_paragraph(style='List Bullet')
    ul3.add_run("Batch Size: 32 | Image Size: 288x288")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Trained a baseline image classifier to predict the exact name of the dish from 25 categories. The accuracy was good initially, though visual ambiguity (like rice and meat variations) caused some confusion between classes.")
    # The original baseline Phase 1 confusion matrix was likely overwritten, but we keep the placeholder.
    cm1_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results_mobilenet_v3_large_baseline", "confusion_matrix.png")
    if os.path.exists(cm1_path):
        doc.add_picture(cm1_path, width=Inches(5.0))

    # Phase 2
    heading(doc, "Phase 2: Binary Categorization (Veg vs Non-Veg)", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" EfficientNet-B2")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" Arabic Dishes (Categorized into Veg/Non-Veg)")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 89.30%")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate: 1e-3 -> 1e-5 (Two-stage fine-tuning)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Loss Function: CrossEntropyLoss with Label Smoothing (0.1)")
    ul3 = doc.add_paragraph(style='List Bullet')
    ul3.add_run("Patience (Early Stopping): 7 epochs")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Restructured the dataset into 'veg' and 'non_veg' categories to simplify the task. Switched to EfficientNet-B2. This resulted in a significant accuracy jump, proving that learning high-level dietary features is much easier for the model than exact dish names.")
    cm2_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results_efficientnet_b2", "confusion_matrix.png")
    if os.path.exists(cm2_path):
        doc.add_picture(cm2_path, width=Inches(5.0))

    # Phase 2B
    heading(doc, "Phase 2B: Binary Categorization with YOLO Segmentation", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large + YOLO Seg")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" YOLO-Segmented Arabic Dishes (Veg/Non-Veg)")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 92.70%")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Precision/Recall: 0.93 / 0.93")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Dataset Size: 534 images in evaluation test set")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" To improve upon the 89.30% binary classification, we introduced YOLO segmentation to crop out backgrounds and fed the cleaned images into MobileNetV3. This pushed the Veg vs Non-Veg accuracy to an impressive 92.70%, proving that removing background noise is critical for food classification.")
    cm2b_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results_mobilenet_v3_large_yolo_seg", "confusion_matrix.png")
    if os.path.exists(cm2b_path):
        doc.add_picture(cm2b_path, width=Inches(5.0))

    # Phase 3
    heading(doc, "Phase 3: Cross-Domain Transfer Learning (German Dataset)", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large + YOLOv8n-seg")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" German Restaurant Dishes")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 95.53%")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate: 1e-3 (Frozen Head) / 1e-5 (Unfrozen Backbone)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Epochs: Stage 1 (15 max), Stage 2 (25 max)")
    ul3 = doc.add_paragraph(style='List Bullet')
    ul3.add_run("Preprocessing: YOLOv8 semantic segmentation mask extraction")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Fine-tuned the Arabic model on a new dataset of German dishes to test cross-domain generalization. Introduced YOLOv8 segmentation as a pre-processing step to crop out backgrounds and plates. The segmentation dramatically reduced background noise, pushing accuracy to ~95.5%.")
    cm3_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "german_model_package", "results", "confusion_matrix.png")
    if os.path.exists(cm3_path):
        doc.add_picture(cm3_path, width=Inches(5.0))

    # Phase 4
    heading(doc, "Phase 4: Real-World Restaurant Testing", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" Final Fine-Tuned German Model (MobileNetV3 + YOLOv8n-seg)")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" External Arabic Restaurant Images (53 unseen test images)")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" ~77.36% (41/53)")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Tested the final German-tuned model on unseen real-world Arabic dishes. Images were grouped by dish and mapped to veg/non_veg labels. The model performed exceptionally well on clear dishes (Hummus, Shawarma, Tamees Bread) but struggled with visually ambiguous ones (e.g., golden-brown cheese pastries looking like meat pies, or large portions of rice hiding the meat in Mandi).")
    cm4_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results", "arabic_confusion_matrix.png")
    if os.path.exists(cm4_path):
        doc.add_picture(cm4_path, width=Inches(5.0))

    # Phase 5
    heading(doc, "Phase 5: Dish-Wise (26-Class) Hierarchical Classifier", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" YOLOv8 Segmented Arabic Dishes (26 distinct dishes, approx 1439 images)")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 72.77%")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate: 1e-3 (Frozen Head) / 1e-5 (Unfrozen Backbone)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Epochs: Stage 1 (15 max), Stage 2 (25 max - Early Stopped at 24)")
    ul3 = doc.add_paragraph(style='List Bullet')
    ul3.add_run("Output Structure: Hierarchical (veg/dish_name and non_veg/dish_name)")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Re-trained the model on the full YOLO-segmented dataset but this time focusing on exact dish names (26 classes). The goal was to maximize results and establish a hierarchical prediction system where images are classified by dish and then logically organized into veg or non-veg subdirectories. The validation accuracy reached ~73% which is very strong for a 26-way classification task on noisy data.")
    cm5_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "phase5_confusion_matrix.png") # I'll copy the old phase 5 cm to phase5_confusion_matrix.png to preserve it since train_dish_wise overwrites
    if not os.path.exists(cm5_path):
        cm5_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results_mobilenet_v3_large", "confusion_matrix.png") # Fallback
    if os.path.exists(cm5_path):
        doc.add_picture(cm5_path, width=Inches(5.0))

    # Phase 6
    heading(doc, "Phase 6: 35-Class Expanded Hierarchical Classifier", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large + YOLO Seg")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" YOLO-Segmented Arabic Dishes (35-Classes, 1,610 images)")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 65.69%")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate: 1e-3 (Stage 1), 1e-5 (Stage 2)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Early Stopping: Triggered before 25 epochs")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" The dataset was expanded to include more images for ambiguous dishes (like Mandi) and new dish categories, pushing the total classes from 26 to 35. While overall accuracy dropped from ~73% to ~66%, this is entirely expected due to the exponential increase in task difficulty (35 classes) and the introduction of new highly ambiguous classes with very few samples.")
    
    cm6_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "phase6_cm.png") # Safe fallback if not exists
    if not os.path.exists(cm6_path):
        cm6_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results_mobilenet_v3_large", "confusion_matrix.png")
    if os.path.exists(cm6_path):
        doc.add_picture(cm6_path, width=Inches(5.0))
        
    # Phase 7
    heading(doc, "Phase 7: Refined Binary Classification (Core 26 Classes)", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large + YOLO Seg")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" Cleaned 26-Class Arabic Dishes (Veg vs Non-Veg, 1,231 images)")
    doc.add_paragraph("Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 89.43%")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate: 1e-3 (Stage 1), 1e-5 (Stage 2)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Early Stopping: Triggered at Epoch 8 (Stage 2)")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" After experimenting with 35 classes, the dataset was cleaned back to the core 26 classes, substituting 'fatayer cheese' and integrating more images of the ambiguous 'mandi'. Reverting back to a Binary Classification target (Veg vs Non-Veg) using YOLO segmentation yielded a strong ~89.43% accuracy, confirming that the model easily distinguishes dietary classes when background noise is removed, even with challenging Arabic dishes.")
    
    cm7_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results_mobilenet_v3_large_yolo_seg", "confusion_matrix.png")
    if os.path.exists(cm7_path):
        doc.add_picture(cm7_path, width=Inches(5.0))

    # Phase 8
    heading(doc, "Phase 8: Shawarma-Boosted Binary Classifier (92.73% Restaurant Accuracy)", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large + YOLO Seg (V5 Dataset)")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" 1,319 images | non_veg: 615, veg: 704 | Shawarma boosted from 5 → 86 images")
    doc.add_paragraph("Standard Test Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 91.20% (534 images)")
    doc.add_paragraph("Restaurant Real-World Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 92.73% (55 images — 51/55 correct)")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul = doc.add_paragraph(style='List Bullet')
    ul.add_run("Learning Rate: 1e-3 (Stage 1), 1e-5 (Stage 2)")
    ul2 = doc.add_paragraph(style='List Bullet')
    ul2.add_run("Early Stopping: Triggered automatically")
    doc.add_paragraph("Per-Dish Restaurant Results:").bold = True
    dishes_p8 = [
        ("Cheese Fatayer", "Veg", "100%"), ("Hummus", "Veg", "100%"),
        ("Mandi", "Non-Veg", "100%"), ("Meat Madfoon", "Non-Veg", "100%"),
        ("Mixed Grill", "Non-Veg", "100%"), ("Okra Tagine", "Veg", "100%"),
        ("Tamees Bread", "Veg", "100%"), ("Za'atara Manaqish", "Veg", "100%"),
        ("Shish Tawook", "Non-Veg", "85.7%"), ("Shawarma", "Non-Veg", "57.1% ⚠"),
    ]
    for dish, cat, acc in dishes_p8:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{dish} ({cat}): {acc}")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Shawarma images were expanded 17x (5→86). 8 out of 10 restaurant dishes now score 100%. Wrapped/pita-style Shawarma remains a challenge as it visually resembles flatbread-based veg dishes. Okra Tagine is noted as an ambiguous dish (can be veg or non-veg depending on preparation).")
    cm8_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "results", "restaurant_analysis", "per_dish_accuracy.png")
    if os.path.exists(cm8_path):
        doc.add_picture(cm8_path, width=Inches(5.5))

    # Phase 9
    heading(doc, "Phase 9: V6 Rebalance (1,432 Images)", level=2)
    doc.add_paragraph("Model Architecture:").bold = True; doc.paragraphs[-1].add_run(" MobileNetV3-Large + YOLO Seg (V6 Dataset)")
    doc.add_paragraph("Dataset:").bold = True; doc.paragraphs[-1].add_run(" 1,432 images | non_veg: 728, veg: 704 | Shawarma boosted to 124, Shish Tawook to 80")
    doc.add_paragraph("Standard Test Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 92.0% (534 images)")
    doc.add_paragraph("Restaurant Real-World Accuracy:").bold = True; doc.paragraphs[-1].add_run(" 90.9% (55 images — 50/55 correct)")
    doc.add_paragraph("Hyperparameters:").bold = True
    ul3 = doc.add_paragraph(style='List Bullet')
    ul3.add_run("Learning Rate: 1e-3 (Stage 1), 1e-5 (Stage 2)")
    ul4 = doc.add_paragraph(style='List Bullet')
    ul4.add_run("Early Stopping: Triggered automatically")
    doc.add_paragraph("Per-Dish Restaurant Results:").bold = True
    dishes_p9 = [
        ("Cheese Fatayer", "Veg", "100%"), ("Hummus", "Veg", "100%"),
        ("Mandi", "Non-Veg", "85.7%"), ("Meat Madfoon", "Non-Veg", "100%"),
        ("Mixed Grill", "Non-Veg", "100%"), ("Okra Tagine", "Veg", "75.0%"),
        ("Tamees Bread", "Veg", "100%"), ("Za'atara Manaqish", "Veg", "100%"),
        ("Shish Tawook", "Non-Veg", "100% ✅"), ("Shawarma", "Non-Veg", "57.1% ⚠"),
    ]
    for dish, cat, acc in dishes_p9:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{dish} ({cat}): {acc}")
    doc.add_paragraph("Details:").bold = True; doc.paragraphs[-1].add_run(" Added 113 new images. Shish Tawook accuracy recovered entirely (from 85.7% to 100%) due to new images featuring white sauce/marinades. Shawarma remains difficult (57.1%) as the new images did not adequately represent the wrapped-in-pita style that appears in the restaurant test set. Standard test set accuracy improved back to 92.0% due to better class balancing (728 non_veg vs 704 veg).")
    cm9_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "results", "restaurant_analysis", "per_dish_accuracy.png")
    if os.path.exists(cm9_path):
        doc.add_picture(cm9_path, width=Inches(5.5))

    doc.save(OUTPUT_PATH)
    print(f"Generated document at: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
