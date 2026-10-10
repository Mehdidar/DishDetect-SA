"""
generate_pipeline_doc.py
Generates the Saudi Arabia Food Classifier Pipeline document
with a professional flowchart and detailed stage descriptions.
"""

import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import matplotlib.patheffects as pe
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT   = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS   = os.path.join(ROOT, "docs")
CHART  = os.path.join(DOCS, "pipeline_flowchart.png")
OUTPUT = os.path.join(DOCS, "Saudi_Arabia_Pipeline.docx")
os.makedirs(DOCS, exist_ok=True)

# ─────────────────────────────────────────────────────────────────────────────
# 1.  FLOWCHART
# ─────────────────────────────────────────────────────────────────────────────
def generate_flowchart():
    fig, ax = plt.subplots(figsize=(12, 20))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 22)
    ax.axis("off")
    fig.patch.set_facecolor("#0D1B2A")
    ax.set_facecolor("#0D1B2A")

    # ── colour palette ──────────────────────────────────────────────────────
    TITLE_BG   = "#1A1A2E"
    STEP_COLORS = [
        "#16213E",  # 1 – Collection
        "#0F3460",  # 2 – Zero-Shot
        "#533483",  # 3 – PIL
        "#E94560",  # 4 – CLIP
        "#F5A623",  # 5 – YOLOv8
        "#2ECC71",  # 6 – MobileNetV3
    ]
    ACCENT  = "#E94560"
    WHITE   = "#FFFFFF"
    SILVER  = "#B0BEC5"
    GOLD    = "#F5A623"

    # ── title ───────────────────────────────────────────────────────────────
    ax.text(5, 21.2, "Saudi Arabia Food Classifier",
            ha="center", va="center", fontsize=20, fontweight="bold",
            color=WHITE, fontfamily="DejaVu Sans")
    ax.text(5, 20.65, "Multi-Stage Image Processing Pipeline",
            ha="center", va="center", fontsize=12,
            color=SILVER, fontfamily="DejaVu Sans")

    # divider
    ax.plot([0.5, 9.5], [20.3, 20.3], color=ACCENT, lw=1.5, alpha=0.7)

    # ── stage data ──────────────────────────────────────────────────────────
    stages = [
        {
            "num":   "01",
            "title": "Image Data Collection",
            "tools": "Requests  •  Selenium  •  Beautiful Soup",
            "desc":  "Automated scraping of food images from\nonline sources across Saudi Arabia",
            "icon":  "🌐",
            "y":     18.8,
            "color": STEP_COLORS[0],
        },
        {
            "num":   "02",
            "title": "Initial Food Classification",
            "tools": "OpenAI  •  Zero-Shot Classification",
            "desc":  "Filter & verify relevant food images;\nreject non-food / low-quality images",
            "icon":  "🤖",
            "y":     15.8,
            "color": STEP_COLORS[1],
        },
        {
            "num":   "03",
            "title": "Image Pre-Processing",
            "tools": "Pillow  (PIL)",
            "desc":  "Resize, normalise, colour-space conversion\nand basic augmentation",
            "icon":  "🖼️",
            "y":     12.8,
            "color": STEP_COLORS[2],
        },
        {
            "num":   "04",
            "title": "Veg / Non-Veg Classification",
            "tools": "CLIP ViT-B/32",
            "desc":  "Zero-shot multimodal classification\nusing vision-language prompts",
            "icon":  "🥗",
            "y":     9.8,
            "color": STEP_COLORS[3],
        },
        {
            "num":   "05",
            "title": "Object Detection & Cropping",
            "tools": "YOLOv8",
            "desc":  "Detect food region, generate segmentation\nmask and crop background",
            "icon":  "📦",
            "y":     6.8,
            "color": STEP_COLORS[4],
        },
        {
            "num":   "06",
            "title": "Final Classification",
            "tools": "MobileNetV3-Large",
            "desc":  "Classify cropped food image as\nVegetarian or Non-Vegetarian",
            "icon":  "✅",
            "y":     3.8,
            "color": STEP_COLORS[5],
        },
    ]

    # ── draw stages ─────────────────────────────────────────────────────────
    BOX_W, BOX_H = 7.4, 2.2
    LEFT = 1.3

    for i, s in enumerate(stages):
        bx, by = LEFT, s["y"] - BOX_H / 2

        # shadow
        shadow = FancyBboxPatch(
            (bx + 0.07, by - 0.07), BOX_W, BOX_H,
            boxstyle="round,pad=0.15", linewidth=0,
            facecolor="#000000", alpha=0.35, zorder=2
        )
        ax.add_patch(shadow)

        # main box
        box = FancyBboxPatch(
            (bx, by), BOX_W, BOX_H,
            boxstyle="round,pad=0.15", linewidth=2,
            edgecolor=s["color"], facecolor=s["color"] + "33", zorder=3
        )
        ax.add_patch(box)

        # left accent bar
        bar = FancyBboxPatch(
            (bx, by), 0.38, BOX_H,
            boxstyle="round,pad=0.0", linewidth=0,
            facecolor=s["color"], zorder=4
        )
        ax.add_patch(bar)

        # step number in bar
        ax.text(bx + 0.19, by + BOX_H / 2, s["num"],
                ha="center", va="center", fontsize=10, fontweight="bold",
                color=WHITE, zorder=5)

        # icon
        ax.text(bx + 0.75, by + BOX_H / 2 + 0.25, s["icon"],
                ha="center", va="center", fontsize=22, zorder=5)

        # title
        ax.text(bx + 1.08, by + BOX_H - 0.38, s["title"],
                ha="left", va="center", fontsize=13, fontweight="bold",
                color=WHITE, zorder=5)

        # tools badge
        tools_x = bx + 1.08
        tools_y = by + BOX_H / 2 - 0.05
        ax.text(tools_x, tools_y, s["tools"],
                ha="left", va="center", fontsize=9.5,
                color=s["color"] if s["color"] != STEP_COLORS[3] else "#FF8FA3",
                fontweight="bold", zorder=5)

        # description
        ax.text(bx + 1.08, by + 0.28, s["desc"],
                ha="left", va="center", fontsize=8.5,
                color=SILVER, zorder=5, linespacing=1.5)

        # connector arrow (skip last)
        if i < len(stages) - 1:
            arrow_x = LEFT + BOX_W / 2
            arrow_y_top = by - 0.04
            arrow_y_bot = stages[i + 1]["y"] + BOX_H / 2 + 0.04
            mid_y = (arrow_y_top + arrow_y_bot) / 2

            ax.annotate("",
                xy=(arrow_x, arrow_y_bot),
                xytext=(arrow_x, arrow_y_top),
                arrowprops=dict(
                    arrowstyle="-|>",
                    color=GOLD,
                    lw=2.2,
                    mutation_scale=18,
                ),
                zorder=6
            )
            # label on arrow
            step_labels = ["filter", "preprocess", "classify", "crop", "classify"]
            ax.text(arrow_x + 0.25, mid_y, step_labels[i],
                    ha="left", va="center", fontsize=7.5,
                    color=GOLD, fontstyle="italic", zorder=7)

    # ── output badge ────────────────────────────────────────────────────────
    out_y = 2.15
    out_box = FancyBboxPatch(
        (2.5, out_y - 0.55), 5.0, 1.0,
        boxstyle="round,pad=0.2", linewidth=2,
        edgecolor=STEP_COLORS[5], facecolor=STEP_COLORS[5] + "55", zorder=3
    )
    ax.add_patch(out_box)
    ax.text(5.0, out_y + 0.12, "OUTPUT",
            ha="center", va="center", fontsize=12, fontweight="bold",
            color=STEP_COLORS[5], zorder=4)
    ax.text(5.0, out_y - 0.26, "Vegetarian  /  Non-Vegetarian  Label",
            ha="center", va="center", fontsize=9.5,
            color=WHITE, zorder=4)

    ax.annotate("",
        xy=(5.0, out_y - 0.55),
        xytext=(5.0, 2.67),
        arrowprops=dict(arrowstyle="-|>", color=STEP_COLORS[5], lw=2.2, mutation_scale=18),
        zorder=6
    )

    # ── footnote ────────────────────────────────────────────────────────────
    ax.text(5, 0.5,
            f"Saudi Arabia Food Classifier Pipeline  •  {datetime.now().strftime('%d %B %Y')}",
            ha="center", va="center", fontsize=8, color=SILVER, alpha=0.7)

    plt.tight_layout(pad=0.3)
    plt.savefig(CHART, dpi=180, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close()
    print(f"  Flowchart saved to: {CHART}")


# ─────────────────────────────────────────────────────────────────────────────
# 2.  WORD DOCUMENT
# ─────────────────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.upper())
    tcPr.append(shd)

def heading(doc, text, level=1, color="1F3864"):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return h

def body(doc, text, italic=False, bold=False, size=11):
    par = doc.add_paragraph()
    r   = par.add_run(text)
    r.italic    = italic
    r.bold      = bold
    r.font.size = Pt(size)
    return par

def bullet(doc, label, value, label_color="1F3864"):
    par = doc.add_paragraph(style="List Bullet")
    r   = par.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(label_color)
    par.add_run("  " + value)
    return par

def stage_table(doc, stages):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = tbl.rows[0].cells
    for i, h in enumerate(["Stage", "Component", "Technology / Tool", "Purpose"]):
        hdr[i].text = h
        set_cell_bg(hdr[i], "1F3864")
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
            run.bold = True

    alt = ["F2F6FC", "FFFFFF"]
    for idx, s in enumerate(stages):
        rc = tbl.add_row().cells
        rc[0].text = s[0]
        rc[1].text = s[1]
        rc[2].text = s[2]
        rc[3].text = s[3]
        for cell in rc:
            set_cell_bg(cell, alt[idx % 2])
    return tbl


def generate_document():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # ── Title ────────────────────────────────────────────────────────────────
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Saudi Arabia Food Classifier\nProject Pipeline")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string("1F3864")

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(
        "Multi-stage pipeline for food image collection, filtering, "
        "cropping and classification\nGenerated: " + datetime.now().strftime("%d %B %Y")
    )
    sr.font.size = Pt(11)
    sr.italic = True
    sr.font.color.rgb = RGBColor.from_string("595959")

    doc.add_page_break()

    # ── 1. Pipeline Overview ─────────────────────────────────────────────────
    heading(doc, "1. Pipeline Overview")
    body(doc,
        "The Saudi Arabia Food Classifier project uses a six-stage automated pipeline "
        "to collect, filter, pre-process, and classify food images sourced from online "
        "platforms. Each stage applies a specialized tool or model to progressively "
        "refine the data until a final Vegetarian / Non-Vegetarian label is produced."
    )

    # ── 2. Flowchart ─────────────────────────────────────────────────────────
    doc.add_paragraph()
    heading(doc, "2. Pipeline Flowchart")
    body(doc, "The diagram below illustrates the end-to-end data flow through all six stages:", size=10)
    doc.add_paragraph()

    if os.path.exists(CHART):
        doc.add_picture(CHART, width=Inches(6.2))
        cp = doc.add_paragraph("Figure 1 — Saudi Arabia Food Classifier: End-to-End Pipeline Flowchart")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.runs[0].italic = True
        cp.runs[0].font.size = Pt(9)
        cp.runs[0].font.color.rgb = RGBColor.from_string("595959")
    else:
        body(doc, "[Flowchart image not found]", italic=True)

    doc.add_page_break()

    # ── 3. Stage Descriptions ─────────────────────────────────────────────────
    heading(doc, "3. Stage-by-Stage Description")

    # Stage 1
    heading(doc, "Stage 1 — Image Data Collection", level=2, color="1F497D")
    bullet(doc, "Tools:", "Requests, Selenium, Beautiful Soup")
    bullet(doc, "Input:", "Search queries / URLs of food image repositories")
    bullet(doc, "Output:", "Raw image files saved to local dataset directory")
    body(doc,
        "Food images are collected programmatically from online sources relevant to the "
        "Saudi Arabian market. Requests handles direct HTTP downloads, Selenium automates "
        "JavaScript-heavy platforms, and Beautiful Soup parses HTML to extract image links. "
        "Images are saved in bulk with metadata (source URL, timestamp, query term)."
    )

    doc.add_paragraph()

    # Stage 2
    heading(doc, "Stage 2 — Initial Food Classification (Zero-Shot)", level=2, color="1F497D")
    bullet(doc, "Tools:", "OpenAI API — Zero-Shot Classification")
    bullet(doc, "Input:", "Raw collected images")
    bullet(doc, "Output:", "Filtered set of confirmed food images")
    body(doc,
        "Not all scraped images are valid food images. OpenAI's zero-shot classification "
        "model is used to verify that each image actually depicts food. Images classified "
        "as non-food (e.g., packaging, advertisements, people) are discarded at this stage, "
        "significantly improving dataset quality before any heavy processing begins."
    )

    doc.add_paragraph()

    # Stage 3
    heading(doc, "Stage 3 — Image Pre-Processing", level=2, color="1F497D")
    bullet(doc, "Tools:", "Pillow (PIL)")
    bullet(doc, "Input:", "Verified food images")
    bullet(doc, "Output:", "Standardized, normalized images ready for model inference")
    body(doc,
        "All verified images pass through a preprocessing pipeline built with Pillow: "
        "resizing to a consistent resolution, RGB colour-space conversion, normalisation "
        "of pixel values, and basic augmentation (flipping, brightness jitter) to improve "
        "downstream model robustness. Images that fail to open or are corrupt are logged "
        "and discarded."
    )

    doc.add_paragraph()

    # Stage 4
    heading(doc, "Stage 4 — Vegetarian / Non-Vegetarian Classification (CLIP)", level=2, color="1F497D")
    bullet(doc, "Tools:", "CLIP ViT-B/32  (OpenAI)")
    bullet(doc, "Input:", "Pre-processed food images")
    bullet(doc, "Output:", "Veg / Non-Veg label with confidence score")
    body(doc,
        "CLIP (Contrastive Language-Image Pre-Training) ViT-B/32 performs zero-shot "
        "vegetarian / non-vegetarian classification using vision-language prompts. "
        "The model compares the image embedding against text prompts such as "
        "'a vegetarian dish' and 'a non-vegetarian dish with meat' to assign a label. "
        "This stage provides a fast, training-free first pass classification."
    )

    doc.add_paragraph()

    # Stage 5
    heading(doc, "Stage 5 — Object Detection & Cropping (YOLOv8)", level=2, color="1F497D")
    bullet(doc, "Tools:", "YOLOv8  (Ultralytics)")
    bullet(doc, "Input:", "Pre-processed food images")
    bullet(doc, "Output:", "Cropped food-region images (background removed)")
    body(doc,
        "YOLOv8 is applied to detect the primary food object within each image and generate "
        "a tight crop, removing plates, tables, cutlery, and background clutter. "
        "Using the segmentation variant (yolov8n-seg), a pixel-level mask is generated "
        "and the food is extracted. This step was the single largest accuracy improvement "
        "in the project — removing background noise forces the downstream CNN to focus "
        "exclusively on food texture and colour."
    )

    doc.add_paragraph()

    # Stage 6
    heading(doc, "Stage 6 — Final Classification (MobileNetV3-Large)", level=2, color="1F497D")
    bullet(doc, "Tools:", "MobileNetV3-Large  (PyTorch / TorchVision)")
    bullet(doc, "Input:", "YOLO-cropped food images")
    bullet(doc, "Output:", "Final Vegetarian / Non-Vegetarian label + confidence score")
    body(doc,
        "The cropped images are passed to a fine-tuned MobileNetV3-Large model — "
        "the primary classifier of the project. The model was trained using two-stage "
        "transfer learning on a dataset of ~1 400 Arabic food images (Veg / Non-Veg labels), "
        "achieving 92.7% accuracy on the standard test set and 90.9% on real-world "
        "Saudi restaurant images. MobileNetV3-Large was chosen for its high accuracy, "
        "low parameter count (~5M), and suitability for deployment on edge devices."
    )

    doc.add_page_break()

    # ── 4. Summary Table ──────────────────────────────────────────────────────
    heading(doc, "4. Pipeline Stage Summary")
    stage_table(doc, [
        ["Stage 1", "Image Data Collection",
         "Requests, Selenium, Beautiful Soup",
         "Automated scraping of food images from online sources"],
        ["Stage 2", "Initial Food Classification",
         "OpenAI Zero-Shot Classification",
         "Filter and verify relevant food images; discard non-food"],
        ["Stage 3", "Image Pre-Processing",
         "Pillow (PIL)",
         "Resize, normalise, colour convert and augment images"],
        ["Stage 4", "Veg / Non-Veg Classification",
         "CLIP ViT-B/32",
         "Zero-shot multimodal dietary classification"],
        ["Stage 5", "Object Detection & Cropping",
         "YOLOv8 (Ultralytics)",
         "Detect food region, generate mask, crop background"],
        ["Stage 6", "Final Classification",
         "MobileNetV3-Large (PyTorch)",
         "Fine-tuned CNN — final Veg / Non-Veg label"],
    ])

    doc.add_page_break()

    # ── 5. Technology Stack ───────────────────────────────────────────────────
    heading(doc, "5. Technology Stack")

    heading(doc, "Data Collection", level=2)
    bullet(doc, "Requests:",    "HTTP library for direct image URL downloads")
    bullet(doc, "Selenium:",    "Browser automation for JavaScript-rendered pages")
    bullet(doc, "BeautifulSoup:", "HTML parser for extracting image links and metadata")

    heading(doc, "AI / ML Models", level=2)
    bullet(doc, "OpenAI Zero-Shot:",  "GPT-based classification — no task-specific training required")
    bullet(doc, "CLIP ViT-B/32:",     "Vision-Language model — classifies using natural language prompts")
    bullet(doc, "YOLOv8:",            "State-of-the-art real-time object detection & segmentation")
    bullet(doc, "MobileNetV3-Large:", "Lightweight CNN fine-tuned for Veg/Non-Veg binary classification")

    heading(doc, "Image Processing", level=2)
    bullet(doc, "Pillow (PIL):", "Image I/O, resizing, colour conversion, normalisation, augmentation")
    bullet(doc, "PyTorch:",      "Deep learning framework for model training and inference")
    bullet(doc, "TorchVision:",  "Pre-trained MobileNetV3-Large backbone and standard transforms")

    doc.add_page_break()

    # ── 6. Design Decisions ──────────────────────────────────────────────────
    heading(doc, "6. Key Design Decisions & Rationale")

    body(doc,
        "The pipeline was designed with a philosophy of progressive refinement — each stage "
        "produces cleaner, more focused data for the next stage, reducing the burden on the "
        "final classification model. Key design decisions include:"
    )
    doc.add_paragraph()

    bullet(doc, "Why CLIP before MobileNetV3?",
           "CLIP provides a fast, zero-shot first pass that filters the most obvious errors "
           "without requiring any labelled training data. MobileNetV3 then performs the "
           "precision fine-grained classification on the already-filtered, YOLO-cropped images.")
    bullet(doc, "Why YOLOv8 segmentation vs simple crop?",
           "Segmentation masks allow pixel-level background removal rather than a rough bounding "
           "box crop. This means the classifier sees only food pixels, not partial plates or tables, "
           "which was the single largest accuracy gain in the project (+3.4pp).")
    bullet(doc, "Why MobileNetV3-Large as the final classifier?",
           "With only ~5M parameters it is significantly lighter than EfficientNet-B2 (~9M) "
           "while achieving higher accuracy on this task (92.7% vs 84.0% in the architecture "
           "comparison experiment) due to full two-stage fine-tuning.")
    bullet(doc, "Why Selenium + Requests + BeautifulSoup together?",
           "Modern websites serve images via JavaScript (requiring Selenium), static HTML "
           "(requiring BeautifulSoup), and direct CDN URLs (requiring Requests). "
           "Using all three tools maximises coverage across diverse source types.")

    doc.add_paragraph()
    body(doc, "Generated: " + datetime.now().strftime("%d %B %Y, %H:%M"), italic=True, size=9)

    doc.save(OUTPUT)
    print(f"  Word document saved to: {OUTPUT}")


# ─────────────────────────────────────────────────────────────────────────────
# ENTRY POINT
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("\nGenerating flowchart...")
    generate_flowchart()
    print("Generating Word document...")
    generate_document()
    print("\nDone!")
